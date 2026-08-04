import sys
import json
import tempfile
from pathlib import Path

# Ensure project root directory is in sys.path
root_dir = Path(__file__).resolve().parent.parent.parent
if str(root_dir) not in sys.path:
    sys.path.insert(0, str(root_dir))

from backend.ai.analyze import analyze_document

SAMPLE_LEGAL_TEXT = """
NON-DISCLOSURE AND CONFIDENTIALITY AGREEMENT

This Non-Disclosure Agreement ("Agreement") is entered into as of January 1, 2026, by and between TechCorp Solutions Pvt. Ltd. ("Disclosing Party") and InnovateX Analytics ("Receiving Party").

1. Purpose & Scope
The Disclosing Party agrees to share proprietary technical specifications, software architecture designs, and artificial intelligence source code with the Receiving Party solely for evaluating a potential commercial joint venture.

2. Important Clauses & Payment Terms
- All disclosures must be marked as "Confidential" or designated as proprietary within 10 business days.
- The evaluation period shall last for 6 months from the effective date.
- Neither party shall be required to make any financial payment under this evaluation phase.

3. Risks & Liabilities
- In the event of an unauthorized disclosure or breach of confidentiality by the Receiving Party, liquidated damages up to INR 5,00,000 may be claimed along with injunctive relief.
- Disclosing Party provides all information "AS IS" without warranties of accuracy or completeness.

4. Obligations
- The Receiving Party agrees to maintain strict confidentiality using the same standard of care used for its own confidential records.
- The Receiving Party shall not reverse engineer, decompile, or create derivative works from the disclosed proprietary technical materials.
"""


def create_sample_pdf(file_path: Path):
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    rect = fitz.Rect(50, 50, 550, 750)
    page.insert_textbox(rect, SAMPLE_LEGAL_TEXT, fontsize=10)
    doc.save(str(file_path))
    doc.close()


def create_sample_docx(file_path: Path):
    from docx import Document
    doc = Document()
    doc.add_heading("NON-DISCLOSURE AGREEMENT", level=1)
    for paragraph in SAMPLE_LEGAL_TEXT.strip().split("\n\n"):
        doc.add_paragraph(paragraph)
    doc.save(str(file_path))


def create_sample_image(file_path: Path):
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (900, 600), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    lines = SAMPLE_LEGAL_TEXT.strip().split("\n")
    y = 20
    for line in lines[:20]:  # Draw first lines onto image canvas
        draw.text((20, y), line, fill=(0, 0, 0))
        y += 24
    img.save(str(file_path))


def test_format(format_name: str, file_path: Path):
    print(f"==================================================")
    print(f" Testing Format: {format_name} ({file_path.name})")
    print(f"==================================================")
    try:
        result = analyze_document(file_path)
        print("Returned JSON Response:")
        print(json.dumps(result, indent=2))
        print("\n")
    except Exception as e:
        print(f"Error testing format {format_name}: {e}\n")


def main():
    temp_dir = Path(tempfile.gettempdir())

    pdf_path = temp_dir / "test_sample_agreement.pdf"
    docx_path = temp_dir / "test_sample_agreement.docx"
    png_path = temp_dir / "test_sample_agreement.png"

    print("Generating sample test files (PDF, DOCX, PNG/JPG)...")
    create_sample_pdf(pdf_path)
    create_sample_docx(docx_path)
    create_sample_image(png_path)

    try:
        test_format("PDF Document", pdf_path)
        test_format("DOCX Document", docx_path)
        test_format("Image Document (PNG/JPG)", png_path)
    finally:
        for path in [pdf_path, docx_path, png_path]:
            if path.exists():
                try:
                    path.unlink()
                except Exception:
                    pass


if __name__ == "__main__":
    main()
